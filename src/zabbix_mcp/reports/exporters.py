"""
Módulo de Exportação de Relatórios em PDF, XLSX, CSV, JSON e Markdown.
Gera documentos profissionais em disco e retorna metadados e caminhos controlados.
"""

import csv
import json
from typing import Any, Dict

import openpyxl
from openpyxl.styles import Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..config import get_settings
from ..security import safe_report_path


def export_to_pdf(report_data: Dict[str, Any], filename: str) -> str:
    """Gera um documento PDF corporativo completo com tabelas, resumo e conclusões."""
    settings = get_settings()
    out_path = safe_report_path(settings.zabbix_reports_dir, filename)

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1A365D"),
        spaceAfter=10
    )
    h2_style = ParagraphStyle(
        "DocHeading2",
        parent=styles["Heading2"],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#2B6CB0"),
        spaceBefore=12,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        "DocBody",
        parent=styles["Normal"],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#2D3748")
    )

    story = []

    # Cabeçalho
    story.append(Paragraph(report_data.get("titulo", "Relatório Zabbix"), title_style))
    story.append(Paragraph(f"<b>Organização:</b> {report_data.get('organizacao')} | <b>Data:</b> {report_data.get('data_geracao')}", body_style))
    story.append(Paragraph(f"<b>Período Analisado:</b> {report_data.get('periodo_analisado')}", body_style))
    story.append(Spacer(1, 12))

    # Resumo Executivo (Tabela de Indicadores Chave)
    story.append(Paragraph("1. Resumo Executivo de Disponibilidade e Alertas", h2_style))
    resumo = report_data.get("resumo_executivo", {})
    resumo_table_data = [
        ["Métrica", "Valor"],
        ["Total de Dispositivos", str(resumo.get("total_hosts", 0))],
        ["Hosts Monitorados", str(resumo.get("hosts_monitorados", 0))],
        ["Hosts Desativados", str(resumo.get("hosts_desativados", 0))],
        ["Total de Problemas no Período", str(resumo.get("total_problemas_registrados", 0))],
        ["Problemas Críticos (Alta/Desastre)", str(resumo.get("problemas_alta_e_desastre", 0))],
        ["Disponibilidade Média Observada", str(resumo.get("disponibilidade_media_observada", "100.0%"))],
    ]
    t_resumo = Table(resumo_table_data, colWidths=[240, 240])
    t_resumo.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2B6CB0")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F7FAFC")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
    ]))
    story.append(t_resumo)
    story.append(Spacer(1, 14))

    # Tabelas de dados
    for tab in report_data.get("tabelas", []):
        story.append(Paragraph(f"Tabela: {tab.get('titulo')}", h2_style))
        raw_rows = [tab.get("cabecalho", [])] + tab.get("linhas", [])
        table_cells = []
        for row in raw_rows:
            table_cells.append([Paragraph(str(cell), body_style) for cell in row])

        t_data = Table(table_cells, colWidths=[120] * len(raw_rows[0]))
        t_data.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4A5568")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ]))
        story.append(t_data)
        story.append(Spacer(1, 12))

    # Conclusões e Recomendações
    story.append(Paragraph("2. Conclusões e Recomendações Técnicas", h2_style))
    for conc in report_data.get("conclusoes", []):
        story.append(Paragraph(f"• {conc}", body_style))
        story.append(Spacer(1, 3))

    doc.build(story)
    return str(out_path)


def export_to_xlsx(report_data: Dict[str, Any], filename: str) -> str:
    """Gera uma planilha Excel (.xlsx) estruturada em múltiplas abas com formatação profissional."""
    settings = get_settings()
    out_path = safe_report_path(settings.zabbix_reports_dir, filename)

    wb = openpyxl.Workbook()
    # Aba 1: Resumo
    ws_resumo = wb.active
    ws_resumo.title = "Resumo Executivo"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws_resumo["A1"] = report_data.get("titulo", "Relatório Zabbix")
    ws_resumo["A1"].font = Font(size=14, bold=True, color="1F497D")

    ws_resumo["A2"] = f"Gerado em: {report_data.get('data_geracao')} | Período: {report_data.get('periodo_analisado')}"
    ws_resumo["A2"].font = Font(size=9, italic=True)

    ws_resumo["A4"] = "Indicador"
    ws_resumo["B4"] = "Valor"
    ws_resumo["A4"].font = header_font
    ws_resumo["A4"].fill = header_fill
    ws_resumo["B4"].font = header_font
    ws_resumo["B4"].fill = header_fill

    resumo = report_data.get("resumo_executivo", {})
    row_idx = 5
    for k, v in resumo.items():
        ws_resumo.cell(row=row_idx, column=1, value=k.replace("_", " ").capitalize()).border = border_thin
        ws_resumo.cell(row=row_idx, column=2, value=str(v)).border = border_thin
        row_idx += 1

    # Abas adicionais para cada tabela
    for tab in report_data.get("tabelas", []):
        tab_title = tab.get("titulo", "Dados")[:30].replace(":", "_").replace("/", "_")
        ws = wb.create_sheet(title=tab_title)

        cabecalho = tab.get("cabecalho", [])
        for col_idx, col_name in enumerate(cabecalho, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill

        for r_idx, row in enumerate(tab.get("linhas", []), 2):
            for c_idx, val in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.border = border_thin

        # Ajusta largura das colunas
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(str(out_path))
    return str(out_path)


def export_to_csv(report_data: Dict[str, Any], filename: str) -> str:
    """Exporta os dados tabulares consolidados para arquivo CSV delimitado por vírgula."""
    settings = get_settings()
    out_path = safe_report_path(settings.zabbix_reports_dir, filename)

    with open(out_path, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["# Relatório Zabbix", report_data.get("titulo")])
        writer.writerow(["# Data Geração", report_data.get("data_geracao")])
        writer.writerow([])

        for tab in report_data.get("tabelas", []):
            writer.writerow([f"## {tab.get('titulo')}"])
            writer.writerow(tab.get("cabecalho", []))
            for row in tab.get("linhas", []):
                writer.writerow(row)
            writer.writerow([])

    return str(out_path)


def export_to_json(report_data: Dict[str, Any], filename: str) -> str:
    """Exporta o relatório completo em formato JSON estruturado."""
    settings = get_settings()
    out_path = safe_report_path(settings.zabbix_reports_dir, filename)

    with open(out_path, mode="w", encoding="utf-8") as f:
        json.dump(report_data, f, ensure_ascii=False, indent=2)

    return str(out_path)


def export_to_markdown(report_data: Dict[str, Any]) -> str:
    """Gera representação em Markdown elegante do relatório para exibição em chat/docs."""
    md = []
    md.append(f"# {report_data.get('titulo')}\n")
    md.append(f"**Data de Geração:** {report_data.get('data_geracao')} | **Período:** {report_data.get('periodo_analisado')}\n")

    md.append("## Resumo Executivo\n")
    resumo = report_data.get("resumo_executivo", {})
    md.append("| Métrica | Valor |")
    md.append("|---|---|")
    for k, v in resumo.items():
        md.append(f"| {k.replace('_', ' ').capitalize()} | {v} |")
    md.append("")

    for tab in report_data.get("tabelas", []):
        md.append(f"### {tab.get('titulo')}\n")
        cab = tab.get("cabecalho", [])
        md.append("| " + " | ".join(cab) + " |")
        md.append("| " + " | ".join(["---"] * len(cab)) + " |")
        for row in tab.get("linhas", []):
            md.append("| " + " | ".join(str(c) for c in row) + " |")
        md.append("")

    md.append("## Conclusões e Recomendações\n")
    for c in report_data.get("conclusoes", []):
        md.append(f"- {c}")

    return "\n".join(md)
